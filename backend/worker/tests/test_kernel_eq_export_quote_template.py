from __future__ import annotations

import sys
from pathlib import Path

# S4（2026-10-19）：报价引擎真源已迁到 backend/kaas_common/export_quote；
# 本文件从 orchestrator 测试迁来，仅重定向 import，断言不变。
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

"""下载模板测试：模板必须「填了就能算」，且与解析器、导出表三边一致。

锁定的事实：
  · 模板列**从产品声明反推**，与解析器共用同一处表头识别 →
    模板表头 ≡ 导出表头 ≡ 解析认得的表头，将来改列名只有一处要改
  · 模板必须覆盖**全部**规格轴与变体包（漏一列整表就都算不了）
  · 反向也要成立：产品输出列里解析器认得、又不是计算列的，不许被漏掉
    （这正是「模板只剩丝径方案/表面处理/工艺、缺了网高与竖丝间距」那个缺陷）
  · 示例值按**表头单位**给：内部 100mm 的竖丝间距，模板里写 10（cm）
  · Excel 默认带下拉框与说明表；CSV 默认只出表头（给脚本生成用）
  · 真正闭环：下载模板 → 填行 → 上传 → 结果与逐条计算一致
"""
import io
import json
from datetime import date
from decimal import Decimal as D

import pytest

from kaas_common.export_quote import QuoteRequest, build
from kaas_common.export_quote.importers import (
    MAX_ROWS,
    ImportError_,
    _build_header_map,
    _norm_header,
    _resolve_column,
    build_template_csv,
    build_template_xlsx,
    parse_table,
    template_columns,
    template_headers,
)
from kaas_common.export_quote.models import Axis, ProductSpec, VariantGroup
from kaas_common.export_quote.pricing import policy_from_dict
from kaas_common.export_quote.products import field_fence

pytestmark = pytest.mark.unit

BASE = "/api/v1/export-quote"


def spec():
    return field_fence.build_spec()


def policy():
    return policy_from_dict({
        "input_vat_rate": D("0.10"),
        "output_vat": {"rate": D("0"), "mode": "additive"},
        "profit": {"mode": "fixed", "value": D("0.1")},
        "refund": {"refund_ratio": D("1"), "give_ratio": D("0.8")},
        # 印花税已并入杂税（misc_fee_rate，对齐平台默认 0.03）
        "misc_fee_rate": D("0.03"),
    })


def req(**kw):
    base = dict(
        product_key="field_fence",
        policy=policy(),
        fx_rate=D("6.7019"),
        orders=field_fence.order_specs((30000,)),
        quote_date=date(2026, 9, 24),
    )
    base.update(kw)
    return QuoteRequest(**base)


# ─────────────────────────────────────────────────────────────────────────────
# 表头识别：单位提取
# ─────────────────────────────────────────────────────────────────────────────
class TestHeaderNormalisation:
    @pytest.mark.parametrize("header,base,unit", [
        ("竖丝间距(cm)", "竖丝间距", "cm"),
        ("Net Height (m)", "netheight", "m"),
        ("mesh_mm", "meshmm", ""),
        ("卷长(m/卷)", "卷长", "m"),          # 括号里带中文量词
        ("卷径(cm)", "卷径", "cm"),
        ("横丝数(根)", "横丝数", ""),          # 纯中文量词不是单位，但括号要去掉
        ("单卷净重(kg)", "单卷净重", "kg"),
        ("网高（m）", "网高", "m"),            # 全角括号
    ])
    def test_unit_is_taken_from_the_header(self, header, base, unit):
        # 基名做过去空格下划线的归一，所以 `net height` / `mesh_mm` 会连起来
        assert _norm_header(header) == (base, unit)

    def test_roll_length_header_is_now_recognised(self):
        """`卷长(m/卷)` 以前因为量词是中文而整列被忽略（静默丢覆盖项）。"""
        p = spec()
        resolved = _resolve_column(p, "卷长(m/卷)", _build_header_map(p))
        assert resolved is not None
        kind, target, _meta, unit = resolved
        # 卷长已是**自由数值**参数（不再是可覆盖的变体属性）
        assert (kind, target, unit) == ("numeric", "roll_len_m", "m")

    def test_computed_columns_are_not_treated_as_inputs(self):
        p = spec()
        maps = _build_header_map(p)
        for header in ("单卷净重(kg)", "单卷售价(元/卷)", "单价(元/米)",
                       "税前成本(元/kg)", "产品"):
            assert _resolve_column(p, header, maps) is None


# ─────────────────────────────────────────────────────────────────────────────
# 模板列：完整性与一致性
# ─────────────────────────────────────────────────────────────────────────────
class TestTemplateColumns:
    def test_covers_every_axis_and_variant(self):
        p = spec()
        cols = template_columns(p)
        assert {c.target for c in cols if c.kind == "variant"} == \
            {v.name for v in p.variants}
        assert {a.name for a in p.axes} <= {c.target for c in cols}
        assert {c.header for c in cols} >= {"网高(m)", "竖丝间距(cm)"}

    def test_every_template_header_is_recognised_by_the_parser(self):
        """模板写出去的每个表头，上传时都必须认得出 —— 同一处识别逻辑。"""
        p = spec()
        maps = _build_header_map(p)
        for c in template_columns(p):
            assert _resolve_column(p, c.header, maps) is not None, c.header

    def test_no_recognised_export_column_is_dropped(self):
        """防漂移：产品输出列里凡是解析器认得的，模板都必须有。

        这条就是那个缺陷的回归测试 —— 模板曾只按 variants 生成，
        于是「网高(m)」「竖丝间距(cm)」两列缺失，下载的模板根本填不出规格。
        「镀锌量」是 spec.coating 的输出列（锌层克数），与「表面处理」
        （变体包标签）语义不同，不进模板是有意为之，故排除。
        """
        p = spec()
        maps = _build_header_map(p)
        excluded = {"镀锌量"}
        should = {c.label for c in p.columns
                  if c.label and _resolve_column(p, c.label, maps) is not None}
        should -= excluded
        assert should <= set(template_headers(p))
        # 尺寸类改为自由数值后，可填列 = 表面处理 + 工艺 + 7 个尺寸类
        # （网高/竖丝间距/边丝径/内横丝径/竖丝径/卷长/卷径）
        assert len(should) >= 9

    def test_required_flags(self):
        cols = {c.header: c for c in template_columns(spec())}
        # 可枚举的必须给：它们决定查价与规格
        assert cols["表面处理"].required
        assert cols["工艺"].required
        # 尺寸类是**自由数值**：产品给了默认值，留空即用默认，所以不是必填。
        # 但模板示例行会写出默认值，用户看得到该按什么量级填。
        for optional in ("网高(m)", "竖丝间距(cm)", "边丝径(mm)",
                         "内横丝径(mm)", "竖丝径(mm)",
                         "卷长(m/卷)", "卷径(cm)"):
            assert not cols[optional].required, optional
            assert cols[optional].default is not None, optional

    def test_examples_use_the_header_unit(self):
        """内部 100mm 的竖丝间距，在 cm 表头下示例值必须是 10。

        自由数值列没有候选值清单（连续量枚举不出来），
        所以"该填什么"靠 `default` 表达，且必须已换算成**表头单位**。
        """
        cols = {c.header: c for c in template_columns(spec())}
        assert cols["竖丝间距(cm)"].options == []
        assert cols["竖丝间距(cm)"].default == 10     # 100mm ÷ 10
        assert cols["网高(m)"].default == 1.2
        assert cols["卷径(cm)"].default == 30        # 0.30m ÷ 0.01

    def test_internal_columns_stay_out(self):
        """内部核算列（成本/利润/税额）绝不进模板 —— 那是给客户看的文件。"""
        p = spec()
        headers = set(template_headers(p))
        assert not (headers & {c.label for c in p.columns if c.internal})
        assert not any(c.startswith("税前") or "税额" in c for c in headers)


# ─────────────────────────────────────────────────────────────────────────────
# Excel 模板
# ─────────────────────────────────────────────────────────────────────────────
def _load(blob: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(blob))


class TestXlsxTemplate:
    def test_sheets_and_order(self):
        wb = _load(build_template_xlsx(spec()))
        # 解析器只读第一个工作表，所以「规格表」必须排第一
        assert wb.sheetnames[0] == "规格表"
        assert {"示例", "填写说明"} <= set(wb.sheetnames)

    def test_header_row_matches_template_headers(self):
        wb = _load(build_template_xlsx(spec()))
        ws = wb["规格表"]
        assert [c.value for c in ws[1]] == template_headers(spec())

    def test_required_columns_have_dropdowns(self):
        p = spec()
        wb = _load(build_template_xlsx(p))
        ws = wb["规格表"]
        headers = [c.value for c in ws[1]]
        dvs = {str(dv.sqref).split(":")[0][0]: dv
               for dv in ws.data_validations.dataValidation}
        for c in template_columns(p):
            if not c.options:
                continue
            letter = ws.cell(row=1, column=headers.index(c.header) + 1).column_letter
            assert letter in dvs, c.header
            assert "|".join(c.options) in dvs[letter].formula1 or \
                all(o in dvs[letter].formula1 for o in c.options)
            # 只引导不硬拦：合法别名（「普锌」）不该被 Excel 挡在外面
            assert dvs[letter].showErrorMessage is False, c.header

    def test_examples_sheet_holds_the_examples(self):
        wb = _load(build_template_xlsx(spec(), examples=3))
        ws = wb["示例"]
        headers = [c.value for c in ws[1]]
        assert headers == template_headers(spec())
        row = {h: ws.cell(row=2, column=i + 1).value
               for i, h in enumerate(headers)}
        # 数字保持数字，不是字符串
        assert row["网高(m)"] == 1.2
        assert row["竖丝间距(cm)"] == 10
        # 自由数值列按表头单位给出示例值
        assert float(row["边丝径(mm)"]) == 2.0
        assert row["卷径(cm)"] == 30
        # 可枚举的给出候选值
        assert row["表面处理"] and row["工艺"]

    def test_spec_sheet_has_no_data_rows_by_default(self):
        """规格表只有表头 —— 示例放在「示例」表，不会被连带算进去。"""
        ws = _load(build_template_xlsx(spec()))["规格表"]
        assert ws.max_row == 1

    def test_instructions_sheet_lists_every_column(self):
        wb = _load(build_template_xlsx(spec()))
        text = "\n".join(
            str(c.value) for row in wb["填写说明"].iter_rows() for c in row
            if c.value is not None
        )
        for c in template_columns(spec()):
            assert c.header in text
        assert "单位规则" in text and "留空规则" in text and "容错规则" in text
        assert str(MAX_ROWS) in text

    def test_long_option_lists_fall_back_to_a_hidden_sheet(self):
        """候选串超过内联上限时改用隐藏表引用，不能让模板生不出来。"""
        long_names = ["规格方案" + "甲" * 100 + str(i) for i in range(3)]
        p = ProductSpec(
            key="long_opts", name="长候选测试",
            axes=[Axis("height_m", (1.2,), "网高(m)", "Height (m)")],
            variants=[VariantGroup("wire", tuple(
                {"code": "c%d" % i, "name": n} for i, n in enumerate(long_names)
            ), "丝径方案", "Wire")],
        )
        wb = _load(build_template_xlsx(p))
        assert "lists" in wb.sheetnames
        assert wb["lists"].sheet_state == "hidden"
        ws = wb["规格表"]
        dvs = list(ws.data_validations.dataValidation)
        # 变体包（丝径方案）应落在 B 列，且候选太长时改用隐藏表引用
        target = next(dv for dv in dvs if str(dv.sqref).startswith("B"))
        assert target.formula1.startswith("=lists!")
        assert not any(str(dv.sqref).startswith("B") and
                       dv.formula1.startswith('"') for dv in dvs)
        # 而且这份模板真的能解析
        csv_text = build_template_csv(p, examples=1)
        parsed = parse_table(csv_text.encode("utf-8"), p, filename="t.csv")
        assert len(parsed.rows) == 1, [i.to_dict() for i in parsed.errors]


# ─────────────────────────────────────────────────────────────────────────────
# CSV 模板
# ─────────────────────────────────────────────────────────────────────────────
class TestCsvTemplate:
    def test_header_only_when_no_examples(self):
        text = build_template_csv(spec(), examples=0)
        assert text.count("\n") == 1
        assert text.splitlines()[0].split(",")[:3] == ["网高(m)", "竖丝间距(cm)", "边丝径(mm)"]

    def test_example_rows_parse_back_cleanly(self):
        p = spec()
        parsed = parse_table(build_template_csv(p, examples=3).encode("utf-8"), p,
                             filename="t.csv")
        assert len(parsed.rows) == 3
        assert not parsed.errors
        assert parsed.ignored_columns == []
        assert parsed.rows[0]["height_m"] == D("1.2")
        assert parsed.rows[0]["mesh_mm"] == D("100")   # 表头 cm=10 → 内部 100mm

    def test_empty_template_says_what_to_do(self):
        p = spec()
        with pytest.raises(ImportError_) as exc:
            parse_table(build_template_csv(p, examples=0).encode("utf-8"), p,
                        filename="t.csv")
        assert "没有数据行" in str(exc.value)


# ─────────────────────────────────────────────────────────────────────────────
# 接口
# ─────────────────────────────────────────────────────────────────────────────

# S4：原 API 级用例依赖已删除的 orchestrator export_quote API 模块；
# worker 侧同一批端点行为由 tests/test_parity_quote.py 覆盖（Q1–Q26）。
