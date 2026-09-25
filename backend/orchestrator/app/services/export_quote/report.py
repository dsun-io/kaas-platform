"""报价输出：内部核算视图 / 客户报价视图 / xlsx / csv。

两种视图来自**同一份数据**的两种投影：
- 内部视图含成本、利润、税额、退税推导（供业务复核）
- 客户视图自动剔除所有 `internal` 列

这是「对外展示不得暴露成本利润率」铁律的落地点 ——
不是靠人工删列，而是靠列的 internal 标记在生成时切分。
"""
from __future__ import annotations

import csv
import io
import os
from decimal import Decimal
from typing import Any, Sequence

from .engine import QuoteResult
from .numeric import to_decimal


def _is_number(value) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _plain(value) -> Any:
    if not _is_number(value):
        return value
    d = to_decimal(value)
    return d


def _number_format(decimals) -> str:
    if decimals is None:
        return "General"
    if decimals <= 0:
        return "0"
    return "0." + "0" * decimals


def build_internal_view(result: QuoteResult) -> dict:
    """内部核算视图（含成本与税链）。"""
    header, body = result.table("internal")
    return {
        "view": "internal",
        "header": header,
        "rows": [[_plain(v) for v in row] for row in body],
        "horizontal_header": [c.label for c in result.columns("internal")],
        "vertical_header": [c.label_en or c.label for c in result.columns("internal")],
        "vertical_data": [
            [d.resolve(c.key) if not _is_number(d.resolve(c.key)) else float(to_decimal(d.resolve(c.key)))
             for d in result.rows]
            for c in result.columns("internal")
        ],
    }


def build_customer_view(result: QuoteResult, english: bool = False) -> dict:
    """客户报价视图（自动剔除 internal 列）。"""
    cols = result.columns("customer")
    header = [(c.label_en or c.label) if english else c.label for c in cols]
    body = []
    for row in result.rows:
        line = []
        for col in cols:
            v = row.resolve(col.key)
            if _is_number(v):
                d = to_decimal(v)
                line.append(float(d) if col.decimals is None else float(round(d, col.decimals)))
            else:
                line.append(v)
        body.append(line)
    return {
        "view": "customer",
        "header": header,
        "rows": body,
        "horizontal_header": header,
        "vertical_header": [(c.label_en or c.label) if english else c.label for c in cols],
        "vertical_data": [
            [d.resolve(c.key) if not _is_number(d.resolve(c.key))
             else float(to_decimal(d.resolve(c.key))) for d in result.rows]
            for c in cols
        ],
    }


def quote_meta(result: QuoteResult) -> dict:
    """报价单元信息（有效期、汇率、术语、档案）。"""
    req = result.request
    return {
        "product_key": result.product.key,
        "product_name": result.product.name,
        "product_name_en": result.product.name_en,
        "quote_date": result.quote_date.isoformat() if result.quote_date else None,
        "valid_until": result.valid_until.isoformat() if result.valid_until else None,
        "validity_days": req.validity_days,
        "fx_rate": str(req.fx_rate),
        "fx_source": req.fx_source,
        "currency": req.currency,
        "quote_currency": req.quote_currency,
        "rate_set_code": result.rate_set_code,
        "term_codes": list(result.term_codes),
        "quote_as": result.quote_as,
        "combo_count": result.combo_count,
        "orders": [
            {
                "name": o.name,
                "label": o.label,
                "label_en": o.label_en,
                "total": o.total,
                "unit": o.unit,
            }
            for o in req.orders
        ],
        "main_term_summary": result.main_term_summary(),
        "diagnostics": list(result.diagnostics),
    }


def write_xlsx(result: QuoteResult, path: str, sheets: Sequence[str] = ("internal", "customer")) -> str:
    """写出 xlsx：内部明细 + 客户报价（可选）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def add(name: str, view: str):
        cols = result.columns(view)
        header, body = result.table(view)
        ws = wb.create_sheet(name[:31])
        for j, label in enumerate(header, start=1):
            c = ws.cell(row=1, column=j, value=label)
            c.font = head_font
            c.fill = head_fill
        for i, line in enumerate(body, start=1):
            for j, value in enumerate(line, start=1):
                cell = ws.cell(row=2 + i - 1, column=j)
                if _is_number(value):
                    cell.value = float(to_decimal(value))
                    if j <= len(cols) and cols[j - 1].decimals is not None:
                        cell.number_format = _number_format(cols[j - 1].decimals)
                else:
                    cell.value = value
        for j, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = col.width or 14
        ws.freeze_panes = ws.cell(row=2, column=1)
        return ws

    for s in sheets:
        if s == "internal" and result.product.internal_columns():
            add("内部明细", "internal")
        elif s == "customer" and result.product.customer_columns():
            add("客户报价", "customer")

    # 参数与说明
    ws = wb.create_sheet("参数与说明")
    r = 1
    ws.cell(row=r, column=1, value="参数").font = head_font
    ws.cell(row=r, column=2, value="取值").font = head_font
    ws.cell(row=r, column=3, value="单位").font = head_font
    ws.cell(row=r, column=4, value="说明").font = head_font
    r += 1
    for item in result.product.params:
        for j, v in enumerate(list(item)[:4], start=1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    r += 1
    for note in list(result.product.notes) + list(result.diagnostics):
        ws.cell(row=r, column=1, value=note)
        r += 1
    for col, w in (("A", 28), ("B", 20), ("C", 12), ("D", 90)):
        ws.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def write_csv(result: QuoteResult, path: str, view: str = "customer") -> str:
    header, body = result.table(view)
    cols = result.columns(view)
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with io.open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for line in body:
            w.writerow([_csv_cell(v, cols[j] if j < len(cols) else None)
                        for j, v in enumerate(line)])
    return path


def _csv_cell(value, col) -> str:
    if value is None:
        return ""
    if not _is_number(value):
        return str(value)
    d = to_decimal(value)
    if col is not None and col.decimals is not None:
        return f"{d:.{col.decimals}f}"
    s = format(d, "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s or "0"
