# -*- coding: utf-8 -*-
"""量合同分页：按行高累计，算出每个分页点与实际页高利用率。"""
from pathlib import Path

import openpyxl

TIM = Path(r"D:\projects\报价制作\Tim")
ws = openpyxl.load_workbook(TIM / "Sales Contract - Tim - 20261012.xlsx").active

last = ws.max_row
hs = {r: (ws.row_dimensions[r].height or 15) for r in range(1, last + 1)}
breaks = sorted(b.id for b in ws.row_breaks.brk)
print("显式分页点(在该行之后断页):", breaks)

# 内容行（有值或非默认高）
print("\n行高明细（>18pt 或分段点）:")
for r in range(1, last + 1):
    marker = "  <== BREAK" if (r in breaks) else ""
    if hs[r] > 18 or marker or ws[f"A{r}"].value:
        v = ws[f"A{r}"].value
        print(f"  r{r:>3} h={hs[r]:>5.1f}  {str(v)[:52]!r}{marker}")

starts = [1] + [b + 1 for b in breaks]
print("\n各页内容高度:")
for i, s in enumerate(starts):
    e = starts[i + 1] - 1 if i + 1 < len(starts) else last
    tot = sum(hs[r] for r in range(s, e + 1))
    print(f"  第{i + 1}页 r{s}-r{e}: {tot:.0f}pt  (预算 600pt, 利用率 {tot / 600 * 100:.0f}%)")
