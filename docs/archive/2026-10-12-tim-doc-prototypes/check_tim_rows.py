# -*- coding: utf-8 -*-
"""确认 CI 末页行序（Note 是否已在签章区上方）并另存新文件名渲染，避开图片缓存。"""
from pathlib import Path

import openpyxl
import pymupdf

TIM = Path(r"D:\projects\报价制作\Tim")
OUT = Path(r"D:\projects\报价制作\.tools\img_tim2")
OUT.mkdir(parents=True, exist_ok=True)

xlsx = TIM / "Commercial Invoice - Buyer - 20261012.xlsx"
ws = openpyxl.load_workbook(xlsx).active
print("== CI 末页行序（A 列 / F 列）==")
for r in range(34, ws.max_row + 1):
    a = ws[f"A{r}"].value
    f = ws[f"F{r}"].value
    h = ws.row_dimensions[r].height
    if a or f:
        print(f"  r{r} h={h}\n      A: {str(a)[:70]}\n      F: {str(f)[:60]}")

for pdf in sorted(TIM.glob("*.pdf")):
    doc = pymupdf.open(pdf)
    tag = pdf.stem.split(" - ")[0][:20].replace(" ", "_").replace("(", "").replace(")", "")
    for i, p in enumerate(doc):
        p.get_pixmap(dpi=150).save(OUT / f"{tag}__v2_p{i + 1}.png")
print("\nrendered ->", OUT)
